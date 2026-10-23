"""
BidScouter - FastAPI Backend (v2)
Adds: application tracking, link health, CPO refund reminders, improved stats.
"""
import json
import logging
import os
import sys
from datetime import datetime, timedelta
from typing import List, Optional

from io import BytesIO
from fastapi import FastAPI, Depends, Query, HTTPException, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import or_, desc, asc, func

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.database import (
    init_db, get_db, Bid, CrawlLog, AppSettings, SessionLocal,
    APPLICATION_STATUSES, CPO_REFUND_STATUSES
)
from backend.scheduler import start_scheduler, stop_scheduler, update_interval
from backend.crawler.engine import run_all_crawlers, check_reminders
from backend.notifier import notify_new_bids, notify_deadline_reminder, notify_cpo_refund
from backend.linkchecker import check_all_links
from backend.document_requirements import get_required_documents_for_bid, get_docs_summary_for_bid

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(name)s %(levelname)s %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler(
            os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                         "logs", "bidscouter.log"),
            encoding="utf-8",
        ),
    ],
)
logger = logging.getLogger("bidscouter.main")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FRONTEND_DIR = os.path.join(BASE_DIR, "frontend")
SETTINGS_PATH = os.path.join(BASE_DIR, "config", "settings.json")

app = FastAPI(title="BidScouter API", version="2.0.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")


@app.on_event("startup")
async def on_startup():
    init_db()
    start_scheduler()
    logger.info("BidScouter v2 started.")


@app.on_event("shutdown")
async def on_shutdown():
    stop_scheduler()


@app.get("/", response_class=HTMLResponse)
async def root():
    with open(os.path.join(FRONTEND_DIR, "index.html"), encoding="utf-8") as f:
        return f.read()


@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    fav_path = os.path.join(FRONTEND_DIR, "favicon.ico")
    if os.path.exists(fav_path):
        return FileResponse(fav_path)
    raise HTTPException(404)


# ─── Pydantic schemas ────────────────────────────────────────────────────────

class BidOut(BaseModel):
    id: str
    title: str
    company: str
    category: str
    bid_number: str
    posted_date: Optional[datetime]
    deadline: Optional[datetime]
    registration_fee: Optional[float]
    registration_fee_currency: str
    bid_security_type: str
    bid_security_amount: Optional[float]
    bid_security_currency: str
    location: str
    contact_name: str
    contact_phone: str
    contact_email: str
    source_url: str
    source_site: str
    source_site_id: str
    is_new: bool
    created_at: Optional[datetime]
    # Tracking
    application_status: str
    application_notes: str
    applied_date: Optional[datetime]
    doc_purchased_date: Optional[datetime]
    cpo_delivered_date: Optional[datetime]
    sample_delivered_date: Optional[datetime]
    result_date: Optional[datetime]
    cpo_refund_collected: bool
    prepared_documents: Optional[str] = "[]"
    required_docs_count: Optional[int] = 11
    prepared_docs_count: Optional[int] = 0
    docs_summary: Optional[str] = ""
    # Link health
    link_status: str
    link_checked_at: Optional[datetime]
    link_http_code: Optional[int]

    class Config:
        from_attributes = True


class ApplicationUpdate(BaseModel):
    application_status: Optional[str] = None
    application_notes: Optional[str] = None
    applied_date: Optional[datetime] = None
    doc_purchased_date: Optional[datetime] = None
    cpo_delivered_date: Optional[datetime] = None
    sample_delivered_date: Optional[datetime] = None
    result_date: Optional[datetime] = None
    cpo_refund_collected: Optional[bool] = None


class DocumentsUpdate(BaseModel):
    prepared_document_ids: List[str]


# ─── Bids API ─────────────────────────────────────────────────────────────────

@app.get("/api/bids", response_model=List[BidOut])
def get_bids(
    db: Session = Depends(get_db),
    search: str = Query(""),
    category: str = Query(""),
    site: str = Query(""),
    app_status: str = Query("", description="Filter by application_status"),
    sort_by: str = Query("deadline"),
    sort_dir: str = Query("asc"),
    days: int = Query(0),
    show_expired: bool = Query(False),
    hide_dead_links: bool = Query(False),
    limit: int = Query(500),
    offset: int = Query(0),
):
    q = db.query(Bid)

    if not show_expired:
        q = q.filter(or_(Bid.deadline == None, Bid.deadline >= datetime.utcnow()))

    if days > 0:
        q = q.filter(Bid.deadline <= datetime.utcnow() + timedelta(days=days))

    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            Bid.title.ilike(like), Bid.company.ilike(like),
            Bid.category.ilike(like), Bid.location.ilike(like),
            Bid.bid_number.ilike(like), Bid.raw_text.ilike(like),
        ))

    if category:
        q = q.filter(Bid.category.ilike(f"%{category}%"))

    if site:
        q = q.filter(Bid.source_site_id == site)

    if app_status:
        q = q.filter(Bid.application_status == app_status)

    if hide_dead_links:
        q = q.filter(Bid.link_status != "dead")

    sort_col_map = {
        "deadline": Bid.deadline,
        "posted_date": Bid.posted_date,
        "company": Bid.company,
        "title": Bid.title,
        "bid_security_amount": Bid.bid_security_amount,
        "registration_fee": Bid.registration_fee,
        "created_at": Bid.created_at,
        "location": Bid.location,
        "application_status": Bid.application_status,
    }
    col = sort_col_map.get(sort_by, Bid.deadline)
    q = q.order_by(asc(col) if sort_dir == "asc" else desc(col))
    return q.offset(offset).limit(limit).all()


@app.get("/api/bids/export/excel")
def export_bids_excel(
    db: Session = Depends(get_db),
    search: str = Query(""),
    category: str = Query(""),
    site: str = Query(""),
    app_status: str = Query(""),
    sort_by: str = Query("deadline"),
    sort_dir: str = Query("asc"),
    days: int = Query(0),
    show_expired: bool = Query(True),
    hide_dead_links: bool = Query(False),
):
    """Export garment bids directly to formatted Excel workbook (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    q = db.query(Bid)
    if not show_expired:
        q = q.filter(or_(Bid.deadline == None, Bid.deadline >= datetime.utcnow()))
    if days > 0:
        q = q.filter(Bid.deadline <= datetime.utcnow() + timedelta(days=days))
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            Bid.title.ilike(like), Bid.company.ilike(like),
            Bid.category.ilike(like), Bid.location.ilike(like),
            Bid.bid_number.ilike(like), Bid.raw_text.ilike(like),
        ))
    if category:
        q = q.filter(Bid.category.ilike(f"%{category}%"))
    if site:
        q = q.filter(Bid.source_site_id == site)
    if app_status:
        q = q.filter(Bid.application_status == app_status)
    if hide_dead_links:
        q = q.filter(Bid.link_status != "dead")

    sort_col_map = {
        "deadline": Bid.deadline,
        "posted_date": Bid.posted_date,
        "company": Bid.company,
        "title": Bid.title,
        "bid_security_amount": Bid.bid_security_amount,
        "registration_fee": Bid.registration_fee,
        "created_at": Bid.created_at,
        "location": Bid.location,
        "application_status": Bid.application_status,
    }
    col = sort_col_map.get(sort_by, Bid.deadline)
    bids = q.order_by(asc(col) if sort_dir == "asc" else desc(col)).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Garment Bids"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Status", "Bid Title", "Ref / Bid No", "Procuring Company / Entity",
        "Category", "Deadline", "Posted Date", "Doc Fee (ETB)",
        "Security Type", "Security Amount (ETB)", "Location / Apply At",
        "Contact Phone", "Contact Email", "Link Health", "Original Portal URL",
        "Source Site", "Milestones & Notes", "Required Documents Checklist"
    ]
    ws.append(headers)

    # Style Header Row
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.row_dimensions[1].height = 28
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    status_labels = {
        "not_applied": "Not Applied",
        "doc_purchased": "Doc Purchased",
        "cpo_delivered": "CPO Delivered",
        "applied": "Applied",
        "sample_delivered": "Sample Delivered",
        "secured": "Secured ✅",
        "failed": "Failed ❌",
        "cancelled": "Cancelled 🚫",
    }

    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for row_idx, b in enumerate(bids, start=2):
        deadline_str = b.deadline.strftime("%Y-%m-%d %H:%M") if b.deadline else "—"
        posted_str = b.posted_date.strftime("%Y-%m-%d %H:%M") if b.posted_date else "—"
        fee_val = float(b.registration_fee) if b.registration_fee else ""
        sec_val = float(b.bid_security_amount) if b.bid_security_amount else ""
        st_label = status_labels.get(b.application_status, b.application_status)

        # Build comprehensive required document checklist string for this bid
        b_docs = get_required_documents_for_bid(b)
        docs_summary_cell = "\n".join([
            f"{'[✓]' if d.get('is_prepared') else '[ ]'} {d['title']}: {d['description']}"
            for d in b_docs
        ])

        row_data = [
            st_label,
            b.title or "—",
            b.bid_number or "—",
            b.company or "—",
            b.category or "Garment",
            deadline_str,
            posted_str,
            fee_val,
            b.bid_security_type or "—",
            sec_val,
            b.location or "—",
            b.contact_phone or "—",
            b.contact_email or "—",
            (b.link_status or "valid").capitalize(),
            b.source_url or "",
            b.source_site or "e-GP",
            b.application_notes or "",
            docs_summary_cell
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 22

        is_alt = (row_idx % 2 == 0)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if is_alt:
                cell.fill = alt_fill

            # Alignments & formatting
            if col_num in [1, 5, 6, 7, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_num in [8, 10]:
                cell.alignment = Alignment(horizontal="right", vertical="top")
                if cell.value != "":
                    cell.number_format = "#,##0.00"
            elif col_num in [2, 4, 11, 17, 18]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif col_num == 15 and b.source_url:
                cell.hyperlink = b.source_url
                cell.font = Font(name="Calibri", size=10, color="2563EB", underline="single")
                cell.alignment = Alignment(horizontal="left", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(13, min(max_len + 3, 50))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Credits & Metadata
    wb.properties.creator = "Enquzer Getachew"
    wb.properties.title = "BidScouter — Ethiopian Garment Tender Tracker"
    wb.properties.description = "Powered by Fashion-Flow 2026. Created by Enquzer Getachew."
    wb.properties.company = "Fashion-Flow 2026"

    footer_row = len(bids) + 3
    footer_cell = ws.cell(row=footer_row, column=1)
    footer_cell.value = "⚡ Powered by Fashion-Flow 2026 • 👤 Created by Enquzer Getachew"
    footer_cell.font = Font(name="Calibri", size=10, italic=True, bold=True, color="1E3A8A")

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"bidscouter_garment_bids_{datetime.utcnow().strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@app.get("/api/bids/stats/summary")
def get_stats(db: Session = Depends(get_db)):
    now = datetime.utcnow()
    total = db.query(Bid).count()
    active = db.query(Bid).filter(
        or_(
            Bid.deadline == None,
            Bid.deadline >= now,
            Bid.application_status.in_(["not_applied", "doc_purchased", "cpo_delivered", "applied", "sample_delivered"])
        )
    ).count()
    new_today = db.query(Bid).filter(Bid.created_at >= now - timedelta(hours=24)).count()
    upcoming_30h = db.query(Bid).filter(
        Bid.deadline != None, Bid.deadline >= now,
        Bid.deadline <= now + timedelta(hours=30),
    ).count()
    unread = db.query(Bid).filter(Bid.is_new == True).count()
    dead_links = db.query(Bid).filter(Bid.link_status == "dead").count()

    # CPO refund pending
    cpo_pending = db.query(Bid).filter(
        Bid.application_status.in_(list(CPO_REFUND_STATUSES)),
        Bid.cpo_refund_collected == False,
        Bid.bid_security_type.ilike("%cpo%"),
    ).count()

    cats = db.query(Bid.category, func.count(Bid.id)).group_by(Bid.category).all()
    sites = db.query(Bid.source_site, func.count(Bid.id)).group_by(Bid.source_site).all()
    statuses = db.query(Bid.application_status, func.count(Bid.id)).group_by(Bid.application_status).all()

    return {
        "total": total,
        "active": active,
        "new_today": new_today,
        "upcoming_30h": upcoming_30h,
        "unread": unread,
        "dead_links": dead_links,
        "cpo_refund_pending": cpo_pending,
        "by_category": [{"category": c, "count": n} for c, n in cats],
        "by_site": [{"site": s, "count": n} for s, n in sites],
        "by_status": [{"status": s, "count": n} for s, n in statuses],
    }


@app.get("/api/bids/{bid_id}", response_model=BidOut)
def get_bid(bid_id: str, db: Session = Depends(get_db)):
    bid = db.query(Bid).filter(Bid.id == bid_id).first()
    if not bid:
        raise HTTPException(404, "Bid not found")
    bid.is_new = False
    db.commit()
    return bid


@app.get("/api/bids/{bid_id}/documents")
def get_bid_documents(bid_id: str, db: Session = Depends(get_db)):
    """Return tailored document requirements for this specific bid request."""
    bid = db.query(Bid).filter(Bid.id == bid_id).first()
    if not bid:
        raise HTTPException(404, "Bid not found")

    docs = get_required_documents_for_bid(bid)
    total_docs = len(docs)
    prepared_docs = sum(1 for d in docs if d.get("is_prepared"))
    mandatory_docs = sum(1 for d in docs if d.get("mandatory"))

    return {
        "bid_id": bid.id,
        "bid_number": bid.bid_number,
        "company": bid.company,
        "title": bid.title,
        "category": bid.category,
        "deadline": bid.deadline,
        "location": bid.location,
        "source_url": bid.source_url,
        "bid_security_amount": bid.bid_security_amount,
        "bid_security_currency": bid.bid_security_currency or "ETB",
        "registration_fee": bid.registration_fee,
        "registration_fee_currency": bid.registration_fee_currency or "ETB",
        "total_docs": total_docs,
        "prepared_docs": prepared_docs,
        "mandatory_docs": mandatory_docs,
        "progress_percent": round((prepared_docs / total_docs * 100)) if total_docs > 0 else 0,
        "documents": docs,
    }


@app.post("/api/bids/{bid_id}/documents")
def update_bid_documents(bid_id: str, payload: DocumentsUpdate, db: Session = Depends(get_db)):
    """Save which required documents have been prepared for this bid."""
    bid = db.query(Bid).filter(Bid.id == bid_id).first()
    if not bid:
        raise HTTPException(404, "Bid not found")

    bid.prepared_documents = json.dumps(payload.prepared_document_ids)
    bid.updated_at = datetime.utcnow()
    db.commit()

    docs = get_required_documents_for_bid(bid, payload.prepared_document_ids)
    total_docs = len(docs)
    prepared_docs = sum(1 for d in docs if d.get("is_prepared"))

    return {
        "success": True,
        "total_docs": total_docs,
        "prepared_docs": prepared_docs,
        "progress_percent": round((prepared_docs / total_docs * 100)) if total_docs > 0 else 0,
        "documents": docs,
    }


@app.patch("/api/bids/{bid_id}/application")
def update_application(bid_id: str, payload: ApplicationUpdate, db: Session = Depends(get_db)):
    """Update application tracking status and dates."""
    bid = db.query(Bid).filter(Bid.id == bid_id).first()
    if not bid:
        raise HTTPException(404, "Bid not found")

    data = payload.model_dump(exclude_none=True)
    for k, v in data.items():
        setattr(bid, k, v)

    # Auto-set timestamps when status changes
    now = datetime.utcnow()
    if payload.application_status == "applied" and not bid.applied_date:
        bid.applied_date = now
    if payload.application_status == "doc_purchased" and not bid.doc_purchased_date:
        bid.doc_purchased_date = now
    if payload.application_status == "cpo_delivered" and not bid.cpo_delivered_date:
        bid.cpo_delivered_date = now
    if payload.application_status == "sample_delivered" and not bid.sample_delivered_date:
        bid.sample_delivered_date = now
    if payload.application_status in ("secured", "failed", "cancelled") and not bid.result_date:
        bid.result_date = now

    bid.updated_at = now
    db.commit()
    return {"ok": True, "status": bid.application_status}


@app.post("/api/bids/{bid_id}/mark_read")
def mark_read(bid_id: str, db: Session = Depends(get_db)):
    bid = db.query(Bid).filter(Bid.id == bid_id).first()
    if bid:
        bid.is_new = False
        db.commit()
    return {"ok": True}


@app.post("/api/bids/mark_all_read")
def mark_all_read(db: Session = Depends(get_db)):
    db.query(Bid).filter(Bid.is_new == True).update({"is_new": False})
    db.commit()
    return {"ok": True}


@app.delete("/api/bids/{bid_id}")
def delete_bid(bid_id: str, db: Session = Depends(get_db)):
    bid = db.query(Bid).filter(Bid.id == bid_id).first()
    if bid:
        db.delete(bid)
        db.commit()
    return {"ok": True}


@app.get("/api/bids/cpo/pending-refunds")
def get_cpo_pending(db: Session = Depends(get_db)):
    """Bids where CPO should be collected (failed/cancelled but not yet collected)."""
    bids = db.query(Bid).filter(
        Bid.application_status.in_(list(CPO_REFUND_STATUSES)),
        Bid.cpo_refund_collected == False,
        or_(Bid.bid_security_type.ilike("%cpo%"), Bid.bid_security_type.ilike("%bond%")),
    ).all()
    return bids


# ─── Crawler control ─────────────────────────────────────────────────────────

@app.post("/api/crawl/now")
async def crawl_now(background_tasks: BackgroundTasks):
    def _run():
        run_all_crawlers(notify_callback=notify_new_bids)
        check_reminders(notify_callback=notify_deadline_reminder)
        check_all_links(limit=30)
    background_tasks.add_task(_run)
    return {"status": "crawl started"}


@app.post("/api/links/check")
async def check_links_now(background_tasks: BackgroundTasks):
    """Trigger link health check."""
    background_tasks.add_task(check_all_links, 100)
    return {"status": "link check started"}


@app.get("/api/crawl/logs")
def get_crawl_logs(db: Session = Depends(get_db), limit: int = 50):
    logs = db.query(CrawlLog).order_by(desc(CrawlLog.started_at)).limit(limit).all()
    return [{
        "id": l.id, "site_id": l.site_id, "site_name": l.site_name,
        "started_at": l.started_at, "finished_at": l.finished_at,
        "status": l.status, "bids_found": l.bids_found,
        "new_bids": l.new_bids, "error_message": l.error_message,
    } for l in logs]


# ─── Settings API ─────────────────────────────────────────────────────────────

@app.get("/api/settings")
def get_settings():
    with open(SETTINGS_PATH, encoding="utf-8") as f:
        data = json.load(f)
    for site in data.get("credentials", {}).get("sites", []):
        if site.get("password"):
            site["password"] = "••••••••"
    return data


@app.put("/api/settings")
def update_settings(payload: dict):
    with open(SETTINGS_PATH, encoding="utf-8") as f:
        current = json.load(f)
    if "credentials" in payload:
        for ns in payload["credentials"].get("sites", []):
            if ns.get("password") == "••••••••":
                for os_ in current.get("credentials", {}).get("sites", []):
                    if os_["id"] == ns["id"]:
                        ns["password"] = os_["password"]
                        break
    current.update(payload)
    with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
        json.dump(current, f, indent=2, ensure_ascii=False)
    new_hours = payload.get("crawler", {}).get("interval_hours")
    if new_hours:
        update_interval(int(new_hours))
    return {"ok": True}


@app.get("/api/meta/statuses")
def get_statuses():
    return {"statuses": APPLICATION_STATUSES}
